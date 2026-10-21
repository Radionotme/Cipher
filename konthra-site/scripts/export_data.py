#!/usr/bin/env python3
"""Export the public-safe Konthra data from the master Excel file to JSON.

Run locally (the spreadsheet holds DM notes and must NOT be committed):
    pip install openpyxl
    python scripts/export_data.py path/to/Konthra_Dramatis_Personae.xlsx
    python scripts/export_data.py path/to/file.xlsx --preview   # local testing only

A character is published only if PublishApproved is Yes AND ReviewStatus is
not "Needs Review". A relationship is published under the same rule, and only
when both ends are published characters.
"""
import argparse, json, sys
from pathlib import Path
from openpyxl import load_workbook

YES = {"yes", "y", "true", "1"}
OUT = Path(__file__).resolve().parent.parent / "site" / "data"


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h else "" for h in next(it)]
    for r in it:
        if any(c not in (None, "") for c in r):
            yield {h: ("" if c is None else str(c).strip()) for h, c in zip(head, r)}


def publishable(r, preview):
    if r.get("ReviewStatus", "").lower() == "needs review":
        return False
    return preview or r.get("PublishApproved", "").lower() in YES


def status_kind(s):
    s = s.lower()
    return "dead" if s.startswith("deceased") else "active" if s.startswith("active") else "unknown"


def split(v):
    return [p.strip() for p in v.replace(",", ";").split(";") if p.strip()] if v else []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--preview", action="store_true",
                    help="ignore PublishApproved (still skips Needs Review). Do not commit the output.")
    a = ap.parse_args()
    wb = load_workbook(a.xlsx, read_only=True, data_only=True)

    chars, dm_notes = [], []
    for r in rows(wb["Characters"]):
        if r.get("DMNotes"):
            dm_notes.append(r["DMNotes"])
        if not publishable(r, a.preview):
            continue
        chars.append({
            "id": r["ID"].lower(),                      # used for #kon-007 deep links
            "name": r.get("PublicDisplayName") or r["Name"],
            "aliases": split(r.get("Aliases", "")),
            "species": r.get("Species", ""),
            "role": r.get("Role", ""),
            "portrait": r.get("Portrait", ""),          # external URL or assets/portraits/x.webp
            "firstSession": r.get("FirstSession", ""),
            "firstLocation": r.get("FirstLocation", ""),
            "currentLocation": r.get("CurrentLocation", ""),
            "status": r.get("Status", ""),
            "statusKind": status_kind(r.get("Status", "")),
            "faction": r.get("Faction", ""),
            "disposition": r.get("Disposition", ""),
            "threatLevel": r.get("ThreatLevel", ""),
            "summary": r.get("PublicSummary", ""),
            "tags": split(r.get("Tags", "")),
            "sessions": split(r.get("Sessions", "")),   # optional column
        })

    ids = {c["id"] for c in chars}
    rels = []
    for r in rows(wb["Relationships"]):
        s, t = r["SourceID"].lower(), r["TargetID"].lower()
        if publishable(r, a.preview) and s in ids and t in ids:
            rels.append({"id": r["RelationshipID"], "source": s, "target": t,
                         "type": r["RelationshipType"],
                         "label": r.get("PublicDescription") or r["RelationshipType"]})

    blob = json.dumps({"characters": chars, "relationships": rels}, ensure_ascii=False)
    leaks = [n for n in dm_notes if len(n) > 20 and n in blob]
    if leaks or "DMNotes" in blob:
        sys.exit(f"ABORT: DM-only text found in public output ({len(leaks)} match/es).")

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "characters.json").write_text(json.dumps(chars, ensure_ascii=False, indent=1), "utf-8")
    (OUT / "relationships.json").write_text(json.dumps(rels, ensure_ascii=False, indent=1), "utf-8")
    print(f"Wrote {len(chars)} characters, {len(rels)} relationships to {OUT}")
    if a.preview:
        print("PREVIEW MODE: PublishApproved was ignored. Do not commit this output.")


if __name__ == "__main__":
    main()
